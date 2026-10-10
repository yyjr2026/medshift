"""导出为 Excel 周排班表。

输出三个工作表，与科室现行周表的结构保持一致：

1. **排班表**——人员×日期矩阵（周一至周日 7 列），格内填岗位／班次名
2. **岗位配置**——岗位清单与每日需求人数
3. **统计**——每人出勤天数、休息天数、夜班次数、加权总分
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import Assignment, DeptConfig, StaffGroup
from ..stats import compute_stats

WEEKDAY_NAMES = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]

GROUP_LABELS = {
    StaffGroup.DIAGNOSIS: "诊断组",
    StaffGroup.TECHNICIAN: "技术组",
    StaffGroup.NURSING: "护理组",
    StaffGroup.OTHER: "其他",
}

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")


def _week_range(assignments: list[Assignment], week_start: dt.date | None) -> tuple[dt.date, dt.date]:
    if week_start is None:
        if not assignments:
            raise ValueError("排班数据为空，无法推断周起始日")
        week_start = min(a.date for a in assignments)
        week_start -= dt.timedelta(days=week_start.weekday())
    return week_start, week_start + dt.timedelta(days=6)


def export_week(
    config: DeptConfig,
    assignments: Iterable[Assignment],
    out_path: str | Path,
    week_start: dt.date | None = None,
) -> Path:
    """导出一周排班表。

    Args:
        config: 科室配置。
        assignments: 排班数据。
        out_path: 输出 .xlsx 路径。
        week_start: 周一日期；留空则按数据中的最早日期推算。

    Returns:
        实际写出的文件路径。
    """
    items = list(assignments)
    start, end = _week_range(items, week_start)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    by_staff_day: dict[tuple[str, dt.date], Assignment] = {
        (a.staff_id, a.date): a for a in items
    }

    wb = Workbook()
    _write_matrix_sheet(wb, config, by_staff_day, start, end)
    _write_position_sheet(wb, config)
    _write_stats_sheet(wb, config, items, start, end)
    wb.save(out_path)
    return out_path


def _cell_text(config: DeptConfig, item: Assignment | None) -> str:
    if item is None:
        return ""
    position = config.position_by_id(item.position_id) if item.position_id else None
    if position is not None:
        return position.name
    shift = config.shift_by_code(item.shift_code)
    return shift.name if shift else item.shift_code


def _write_matrix_sheet(
    wb: Workbook,
    config: DeptConfig,
    by_staff_day: dict[tuple[str, dt.date], Assignment],
    start: dt.date,
    end: dt.date,
) -> None:
    ws = wb.active
    ws.title = "排班表"

    days = [start + dt.timedelta(days=i) for i in range(7)]

    ws.cell(row=1, column=1, value=f"{config.name}排班表")
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws.cell(row=2, column=1, value=f"{start.isoformat()} 至 {end.isoformat()}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    ws.cell(row=3, column=1, value="姓名")
    ws.cell(row=3, column=2, value="组别")
    for idx, day in enumerate(days):
        col = 3 + idx
        label = f"{WEEKDAY_NAMES[idx]}\n{day.isoformat()[5:]}"
        holiday = config.calendar.label(day)
        if holiday:
            label = f"{label}\n{holiday}"
        ws.cell(row=3, column=col, value=label)

    for col in range(1, 10):
        cell = ws.cell(row=3, column=col)
        cell.font = Font(bold=True)
        cell.alignment = _CENTER
        cell.border = _BORDER
        cell.fill = _HEADER_FILL

    ordered = [
        s for s in config.staff if s.is_rosterable
    ] + [s for s in config.staff if not s.is_rosterable]

    row = 4
    for person in ordered:
        label = person.name
        if person.supervisor_id:
            supervisor = config.staff_by_id(person.supervisor_id)
            if supervisor is not None:
                label = f"{person.name}（随 {supervisor.name}）"
        name_cell = ws.cell(row=row, column=1, value=label)
        name_cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=row, column=2, value=GROUP_LABELS.get(person.group, person.group.value))

        for idx, day in enumerate(days):
            item = by_staff_day.get((person.id, day))
            cell = ws.cell(row=row, column=3 + idx, value=_cell_text(config, item))
            cell.alignment = _CENTER
            if item is not None:
                shift = config.shift_by_code(item.shift_code)
                if shift and shift.color:
                    cell.fill = PatternFill("solid", fgColor=shift.color)

        if not person.is_rosterable:
            note = person.note or person.status.value
            ws.cell(row=row, column=3, value=f"（{note}）")
            ws.cell(row=row, column=3).alignment = _CENTER

        for col in range(1, 10):
            ws.cell(row=row, column=col).border = _BORDER
        row += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    for idx in range(7):
        ws.column_dimensions[get_column_letter(3 + idx)].width = 14
    ws.freeze_panes = "C4"


def _write_position_sheet(wb: Workbook, config: DeptConfig) -> None:
    ws = wb.create_sheet("岗位配置")
    headers = ["岗位代码", "岗位名称", "班次", "每日人数", "组别", "所需技能"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER

    for position in config.positions:
        shift = config.shift_by_code(position.shift_code)
        ws.append(
            [
                position.id,
                position.name,
                shift.name if shift else position.shift_code,
                position.daily_need,
                GROUP_LABELS.get(position.group, "") if position.group else "不限",
                "、".join(position.required_skills),
            ]
        )

    for col, width in zip("ABCDEF", (16, 16, 14, 10, 12, 24)):
        ws.column_dimensions[col].width = width


def _write_stats_sheet(
    wb: Workbook,
    config: DeptConfig,
    items: list[Assignment],
    start: dt.date,
    end: dt.date,
) -> None:
    ws = wb.create_sheet("统计")
    headers = ["姓名", "组别", "出勤天数", "休息天数", "夜班次数", "值出天数", "加权总分"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER

    for row in compute_stats(config, items):
        ws.append(
            [
                row.name,
                row.group,
                row.work_days,
                row.off_days,
                row.night_count,
                row.post_night_count,
                row.weighted_score,
            ]
        )

    for col, width in zip("ABCDEFG", (14, 12, 10, 10, 10, 10, 12)):
        ws.column_dimensions[col].width = width
