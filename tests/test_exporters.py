"""导入导出与统计测试。"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from openpyxl import load_workbook

from medshift import Scheduler
from medshift.errors import AssignmentFormatError
from medshift.exporters import read_assignments, write_assignments
from medshift.models import Assignment
from medshift.stats import compute_stats

WEEK_START = dt.date(2026, 10, 12)


def test_csv_round_trip(tmp_path: Path, assignments):
    out = tmp_path / "round.csv"
    write_assignments(assignments, out)
    again = read_assignments(out)

    assert len(again) == len(assignments)
    assert {a.key() for a in again} == {a.key() for a in assignments}
    assert {a.shift_code for a in again} == {a.shift_code for a in assignments}


def test_csv_missing_file(tmp_path: Path):
    with pytest.raises(AssignmentFormatError):
        read_assignments(tmp_path / "nope.csv")


def test_csv_bad_header(tmp_path: Path):
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    with pytest.raises(AssignmentFormatError):
        read_assignments(bad)


def test_csv_bad_date(tmp_path: Path):
    bad = tmp_path / "baddate.csv"
    bad.write_text("staff_id,date,shift_code\nD01,2026/10/12,OFF\n", encoding="utf-8")
    with pytest.raises(AssignmentFormatError):
        read_assignments(bad)


def test_csv_tags_round_trip(tmp_path: Path):
    item = Assignment(
        staff_id="D01",
        date=WEEK_START,
        shift_code="OFF",
        tags=["backup", "makeup"],
    )
    out = tmp_path / "tags.csv"
    write_assignments([item], out)
    again = read_assignments(out)
    assert again[0].tags == ["backup", "makeup"]


def test_export_excel_structure(tmp_path: Path, config, assignments):
    out = tmp_path / "week.xlsx"
    path = Scheduler(config).export_excel(assignments, out, WEEK_START)
    assert path.exists()

    wb = load_workbook(path)
    assert wb.sheetnames == ["排班表", "岗位配置", "统计"]

    sheet = wb["排班表"]
    assert sheet.cell(row=3, column=1).value == "姓名"
    assert sheet.cell(row=3, column=3).value.startswith("星期一")
    # 表头 7 天 + 姓名 + 组别
    assert sheet.max_column == 9

    # 人员行数 = 可排班 + 进修 + 离岗
    assert sheet.max_row == 3 + len(config.staff)

    stats_sheet = wb["统计"]
    assert stats_sheet.cell(row=1, column=1).value == "姓名"
    assert stats_sheet.max_row == 1 + len(config.assignable_staff())


def test_export_without_week_start(tmp_path: Path, config, assignments):
    out = tmp_path / "auto.xlsx"
    path = Scheduler(config).export_excel(assignments, out)
    assert path.exists()
    assert load_workbook(path)["排班表"].cell(row=2, column=1).value == "2026-10-12 至 2026-10-18"


def test_export_requires_data(tmp_path: Path, config):
    with pytest.raises(ValueError):
        Scheduler(config).export_excel([], tmp_path / "empty.xlsx")


def test_stats_are_consistent(config, assignments):
    rows = compute_stats(config, assignments)
    assert len(rows) == len(config.assignable_staff())
    for row in rows:
        assert row.total_days == 7, f"{row.name} 统计天数不等于 7"
        assert row.work_days >= 0 and row.night_count >= 0
        assert row.night_count <= row.work_days + row.post_night_count

    # 每天 3 个值班岗（诊断 / CT / DR）
    assert sum(r.night_count for r in rows) == 3 * 7


def test_trainees_excluded_from_stats(config, assignments):
    rows = compute_stats(config, assignments)
    names = {r.staff_id for r in rows}
    assert "D99" not in names  # 进修医师
    assert "D10" not in names  # 派驻他科
